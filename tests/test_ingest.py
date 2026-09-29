"""Tests that every supported file type is read correctly and bad files are skipped."""
import csv
from email.message import EmailMessage
from pathlib import Path

import pytest

from src.ingest import ingest_enron_csv, ingest_folder

MARKER = "Quarterly roadmap marker 42"


@pytest.fixture()
def sample_folder(tmp_path: Path) -> Path:
    import docx
    from openpyxl import Workbook
    from reportlab.pdfgen import canvas

    (tmp_path / "note.txt").write_text(f"{MARKER} in a text file", encoding="utf-8")

    msg = EmailMessage()
    msg["From"] = "Asha Rao <asha.rao@example.in>"
    msg["To"] = "team@example.in, lead@example.in"
    msg["Subject"] = "Status update"
    msg["Date"] = "Tue, 14 Mar 2023 10:30:00 +0530"
    msg.set_content(f"{MARKER} in an email body")
    (tmp_path / "mail.eml").write_bytes(bytes(msg))

    with (tmp_path / "sheet.csv").open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows([["Name", "Account No"], ["Test User", "12345678901"]])

    wb = Workbook()
    wb.active.append(["Metric", "Value"])
    wb.active.append([MARKER, 7])
    wb.save(tmp_path / "book.xlsx")

    d = docx.Document()
    d.add_paragraph(f"{MARKER} in a Word file")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "IFSC", "HDFC0001234"
    d.save(tmp_path / "letter.docx")

    sub = tmp_path / "nested" / "deeper"
    sub.mkdir(parents=True)
    c = canvas.Canvas(str(sub / "report.pdf"))
    c.drawString(72, 720, f"{MARKER} in a PDF")
    c.save()

    (tmp_path / "broken.pdf").write_bytes(b"%PDF-1.4 not really a pdf \x00\xff")
    (tmp_path / "broken.xlsx").write_bytes(b"this is not a zip file")
    (tmp_path / "ignored.jpg").write_bytes(b"\xff\xd8")  # unsupported extension: silently ignored
    return tmp_path


def by_type(records):
    return {r["file_type"]: r for r in records}


def test_each_file_type_is_read(sample_folder):
    records, skipped = ingest_folder(sample_folder)
    recs = by_type(records)
    assert set(recs) == {"txt", "eml", "csv", "xlsx", "docx", "pdf"}
    for ft in ["txt", "eml", "xlsx", "docx", "pdf"]:
        assert MARKER in recs[ft]["text"], ft
    assert all(r["char_count"] == len(r["text"]) for r in records)
    assert all(r["doc_id"].startswith("DOC-") for r in records)


def test_email_fields_parsed(sample_folder):
    eml = by_type(ingest_folder(sample_folder)[0])["eml"]
    assert eml["sender"] == "asha.rao@example.in"
    assert "team@example.in" in eml["recipients"] and "lead@example.in" in eml["recipients"]
    assert eml["subject"] == "Status update"
    assert eml["date"].startswith("2023-03-14T05:00:00")
    assert eml["text"].startswith("Status update")


def test_spreadsheet_rows_keep_column_context(sample_folder):
    csv_rec = by_type(ingest_folder(sample_folder)[0])["csv"]
    assert "Account No: 12345678901" in csv_rec["text"]
    docx_rec = by_type(ingest_folder(sample_folder)[0])["docx"]
    assert "IFSC: HDFC0001234" in docx_rec["text"]


def test_corrupt_files_are_skipped_not_fatal(sample_folder):
    records, skipped = ingest_folder(sample_folder)
    names = {Path(s["path"]).name for s in skipped}
    assert names == {"broken.pdf", "broken.xlsx"}
    assert all(s["error"] for s in skipped)
    assert len(records) == 6


def test_truncation(sample_folder, monkeypatch):
    from src import load_config
    cfg = {**load_config(), "ingest": {**load_config()["ingest"], "max_text_length": 10}}
    records, _ = ingest_folder(sample_folder, cfg=cfg)
    assert all(r["char_count"] <= 10 for r in records)
    assert any(r["truncated"] for r in records)


def test_enron_loader_and_limit(tmp_path):
    path = tmp_path / "emails.csv"
    rows = []
    for i in range(5):
        rows.append({"file": f"allen-p/_sent_mail/{i}.", "message": (
            f"Message-ID: <{i}@thyme>\nDate: Mon, 14 May 2001 16:39:00 -0700 (PDT)\n"
            f"From: phillip.allen@enron.com\nTo: tim.belden@enron.com\nSubject: Forecast {i}\n"
            f"Mime-Version: 1.0\nContent-Type: text/plain; charset=us-ascii\n\nHere is our forecast number {i}\n")})
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["file", "message"])
        w.writeheader()
        w.writerows(rows)
    records, skipped = ingest_enron_csv(path, limit=3)
    assert len(records) == 3 and not skipped
    r = records[0]
    assert r["sender"] == "phillip.allen@enron.com"
    assert r["recipients"] == "tim.belden@enron.com"
    assert r["subject"] == "Forecast 0"
    assert "forecast number 0" in r["text"]
    assert r["date"].startswith("2001-05-14T23:39:00")


def test_enron_wrong_columns(tmp_path):
    path = tmp_path / "emails.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(ValueError):
        ingest_enron_csv(path)
