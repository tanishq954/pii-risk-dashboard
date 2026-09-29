"""Document ingestion: turn a folder of mixed files (or the Enron CSV) into plain text.

Every loader returns the same record shape so the detector never needs to
know where a document came from:

    doc_id, source_path, file_type, sender, recipients, date, subject,
    text, char_count, truncated

Design rules:
* **Never crash on a bad file.** Each file is read inside try/except; a
  failure is logged and added to the ``skipped`` list, and the walk continues.
* **Spreadsheets keep their column context.** Each row is written as
  ``Header: value | Header: value`` so a number sits right next to the
  word that explains it ("Account No: ...").  Without this, a bank account
  in row 150 would be 150 lines away from its column header.
* **Long texts are truncated** to ``ingest.max_text_length`` and flagged.
"""
from __future__ import annotations

import csv
import email
import io
import re
from datetime import datetime, timezone
from email import policy
from email.utils import getaddresses, parsedate_to_datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from . import get_logger, load_config

log = get_logger("ingest")

try:  # Outlook .msg support is optional
    import extract_msg  # type: ignore
    HAS_MSG = True
except Exception:  # pragma: no cover - depends on environment
    HAS_MSG = False

csv.field_size_limit(10_000_000)  # Enron messages can be very long


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def _strip_html(html: str) -> str:
    """Very small HTML-to-text conversion (enough for e-mail bodies)."""
    html = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>", "\n", html)
    text = re.sub(r"<[^>]+>", " ", html)
    return re.sub(r"[ \t]+", " ", text)


def _iso(dt: Optional[datetime]) -> Optional[str]:
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc).isoformat(timespec="seconds")


def _mtime(path: Path) -> Optional[str]:
    try:
        return _iso(datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc))
    except OSError:
        return None


def _rows_to_text(header: List[str], rows: List[List[Any]]) -> str:
    """Serialise table rows as ``Header: value | Header: value`` lines."""
    header = [str(h).strip() if h not in (None, "") else f"col{i + 1}" for i, h in enumerate(header)]
    lines = []
    for row in rows:
        cells = [f"{header[i] if i < len(header) else f'col{i + 1}'}: {v}"
                 for i, v in enumerate(row) if v not in (None, "")]
        if cells:
            lines.append(" | ".join(str(c) for c in cells))
    return "\n".join(lines)


def _join_subject(subject: str, body: str) -> str:
    """Emails are analysed as subject + blank line + body (subject can hold PII too)."""
    subject = (subject or "").strip()
    return f"{subject}\n\n{body}" if subject else body


# ---------------------------------------------------------------------------
# Per-format readers. Each returns a partial record dict.
# ---------------------------------------------------------------------------
def parse_email_message(msg: email.message.Message) -> Dict[str, Any]:
    """Split a parsed e-mail into sender, recipients, date, subject and body."""
    sender = str(msg.get("From", "") or "").strip()
    recipients = ", ".join(a for _, a in getaddresses(
        [str(msg.get(h, "")) for h in ("To", "Cc", "Bcc") if msg.get(h)]) if a)
    date = None
    if msg.get("Date"):
        try:
            date = _iso(parsedate_to_datetime(str(msg["Date"])))
        except (TypeError, ValueError, IndexError):
            date = None
    subject = str(msg.get("Subject", "") or "").strip()

    body = ""
    if msg.is_multipart():
        parts = [p for p in msg.walk() if not p.is_multipart() and p.get_content_disposition() != "attachment"]
        plain = [p for p in parts if p.get_content_type() == "text/plain"]
        html = [p for p in parts if p.get_content_type() == "text/html"]
        chosen = plain or html
        texts = []
        for p in chosen:
            payload = p.get_payload(decode=True) or b""
            charset = p.get_content_charset() or "utf-8"
            t = payload.decode(charset, errors="replace")
            texts.append(_strip_html(t) if p.get_content_type() == "text/html" else t)
        body = "\n".join(texts)
    else:
        payload = msg.get_payload(decode=True)
        if isinstance(payload, bytes):
            body = payload.decode(msg.get_content_charset() or "utf-8", errors="replace")
        else:
            body = str(msg.get_payload())
        if msg.get_content_type() == "text/html":
            body = _strip_html(body)

    _, sender_addr = getaddresses([sender])[0] if sender else ("", "")
    return {"sender": sender_addr or sender, "recipients": recipients, "date": date,
            "subject": subject, "text": _join_subject(subject, body)}


def read_eml(path: Path) -> Dict[str, Any]:
    with path.open("rb") as fh:
        msg = email.message_from_binary_file(fh, policy=policy.compat32)
    rec = parse_email_message(msg)
    rec["date"] = rec["date"] or _mtime(path)
    return rec


def read_msg(path: Path) -> Dict[str, Any]:
    if not HAS_MSG:
        raise RuntimeError("extract-msg library not installed; .msg file skipped")
    m = extract_msg.Message(str(path))
    try:
        date = None
        if m.date:
            date = _iso(m.date) if isinstance(m.date, datetime) else _iso(parsedate_to_datetime(str(m.date)))
        subject = m.subject or ""
        return {"sender": m.sender or "", "recipients": m.to or "", "date": date or _mtime(path),
                "subject": subject, "text": _join_subject(subject, m.body or "")}
    finally:
        m.close()


def read_txt(path: Path) -> Dict[str, Any]:
    return {"text": path.read_text(encoding="utf-8", errors="replace")}


def read_csv_file(path: Path, max_rows: int) -> Dict[str, Any]:
    raw = path.read_bytes().decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(raw))
    rows = []
    for i, row in enumerate(reader):
        if i > max_rows:
            break
        rows.append(row)
    if not rows:
        return {"text": ""}
    return {"text": _rows_to_text(rows[0], rows[1:])}


def read_xlsx(path: Path, max_rows: int) -> Dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        parts = []
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > max_rows:
                    break
                rows.append(list(row))
            rows = [r for r in rows if any(v not in (None, "") for v in r)]
            if rows:
                parts.append(f"Sheet: {ws.title}\n" + _rows_to_text(rows[0], rows[1:]))
        return {"text": "\n\n".join(parts)}
    finally:
        wb.close()


def read_docx(path: Path) -> Dict[str, Any]:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        rows = [[c.text.strip() for c in r.cells] for r in table.rows]
        if rows:
            # Two-column tables are usually "label | value" pairs; keep them on one line.
            if all(len(r) == 2 for r in rows):
                parts.extend(f"{a}: {b}" for a, b in rows)
            else:
                parts.append(_rows_to_text(rows[0], rows[1:]))
    created = d.core_properties.created
    return {"text": "\n".join(parts), "date": _iso(created) if created else None}


def read_pdf(path: Path) -> Dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    date = None
    try:
        meta = reader.metadata
        if meta and meta.creation_date:
            date = _iso(meta.creation_date)
    except Exception:
        date = None
    return {"text": text, "date": date}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------
def _finalise(rec: Dict[str, Any], doc_id: str, max_len: int) -> Dict[str, Any]:
    """Fill defaults, truncate long text and compute char_count."""
    text = rec.get("text") or ""
    truncated = len(text) > max_len
    if truncated:
        text = text[:max_len]
    out = {
        "doc_id": doc_id,
        "source_path": rec.get("source_path", ""),
        "file_type": rec.get("file_type", ""),
        "sender": rec.get("sender") or "",
        "recipients": rec.get("recipients") or "",
        "date": rec.get("date"),
        "subject": rec.get("subject") or "",
        "text": text,
        "char_count": len(text),
        "truncated": truncated,
    }
    if truncated:
        log.info("Truncated %s to %d characters", out["source_path"], max_len)
    return out


def ingest_folder(folder: str | Path, limit: Optional[int] = None,
                  cfg: Optional[dict] = None) -> Tuple[List[Dict[str, Any]], List[Dict[str, str]]]:
    """Recursively read every supported file under ``folder``.

    Returns ``(records, skipped)`` where ``skipped`` is a list of
    ``{"path", "error"}`` dicts for files that could not be read.
    """
    cfg = cfg or load_config()
    icfg = cfg["ingest"]
    folder = Path(folder)
    if not folder.exists():
        raise FileNotFoundError(f"Folder not found: {folder}")
    exts = {e.lower() for e in icfg["supported_extensions"]}
    excluded = set(icfg["exclude_file_names"])
    max_len, max_rows = icfg["max_text_length"], icfg["max_spreadsheet_rows"]

    readers = {
        ".eml": read_eml, ".msg": read_msg, ".txt": read_txt, ".docx": read_docx, ".pdf": read_pdf,
        ".csv": lambda p: read_csv_file(p, max_rows), ".xlsx": lambda p: read_xlsx(p, max_rows),
    }
    files = sorted(p for p in folder.rglob("*")
                   if p.is_file() and p.suffix.lower() in exts and p.name not in excluded
                   and not p.name.startswith("~$"))
    if limit:
        files = files[:limit]

    records, skipped = [], []
    for path in files:
        ext = path.suffix.lower()
        try:
            rec = readers[ext](path)
            rec["source_path"] = str(path)
            rec["file_type"] = ext.lstrip(".")
            rec["date"] = rec.get("date") or _mtime(path)
            records.append(_finalise(rec, f"DOC-{len(records) + 1:06d}", max_len))
        except Exception as exc:  # never let one bad file stop the scan
            log.warning("Skipped %s: %s", path.name, exc)
            skipped.append({"path": str(path), "error": f"{type(exc).__name__}: {exc}"[:500]})
    log.info("Ingested %d files from %s (%d skipped)", len(records), folder, len(skipped))
    return records, skipped


def ingest_enron_csv(csv_path: str | Path, limit: Optional[int] = None,
                     cfg: Optional[dict] = None) -> Tuple[List[Dict[str, Any]], List[Dict[str, str]]]:
    """Load Kaggle's Enron ``emails.csv`` (columns ``file``, ``message``).

    The full file holds ~500,000 messages, so only the first ``limit``
    rows are read (default from config: 5000). The file is streamed row by
    row, so memory use stays small whatever the limit.
    """
    cfg = cfg or load_config()
    limit = limit or cfg["ingest"]["enron_default_limit"]
    max_len = cfg["ingest"]["max_text_length"]
    csv_path = Path(csv_path)
    if not csv_path.exists():
        raise FileNotFoundError(
            f"Enron file not found: {csv_path}. Download 'Enron Email Dataset' (wcukierski) from Kaggle "
            f"and place emails.csv in data/raw/.")

    records, skipped = [], []
    with csv_path.open("r", encoding="utf-8", errors="replace", newline="") as fh:
        reader = csv.DictReader(fh)
        if not reader.fieldnames or not {"file", "message"} <= set(reader.fieldnames):
            raise ValueError(f"{csv_path.name} does not look like the Kaggle Enron file "
                             f"(expected columns 'file' and 'message', got {reader.fieldnames})")
        for row in reader:
            if len(records) + len(skipped) >= limit:
                break
            try:
                msg = email.message_from_string(row["message"] or "", policy=policy.compat32)
                rec = parse_email_message(msg)
                rec["source_path"] = f"enron/{row['file']}"
                rec["file_type"] = "enron_email"
                records.append(_finalise(rec, f"DOC-{len(records) + 1:06d}", max_len))
            except Exception as exc:
                skipped.append({"path": f"enron/{row.get('file', '?')}", "error": f"{type(exc).__name__}: {exc}"[:500]})
    log.info("Loaded %d Enron emails (%d skipped)", len(records), len(skipped))
    return records, skipped
