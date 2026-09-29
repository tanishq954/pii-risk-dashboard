"""Generate a synthetic Indian corporate dataset for the PII Risk Dashboard.

*** ALL DATA PRODUCED BY THIS SCRIPT IS FAKE. ***
Names, Aadhaar numbers, PANs, bank accounts, phone numbers and salaries are
generated with Faker (en_IN) and random numbers. Aadhaar numbers pass the
Verhoeff checksum and PAN/IFSC follow the official formats so that the
detectors can be tested realistically, but none belong to real people.

Output (default data/sample/):
* ~300 files: .eml, .txt, .csv, .xlsx, .docx, .pdf
* ground_truth.csv - which PII types were planted in each file, and the
  expected risk tier, so the scanner's precision/recall can be measured.
* one deliberately corrupt PDF, to demonstrate that bad files are skipped.

Mix: ~25% high-risk, ~30% medium-risk, ~45% low-risk or clean.
Fixed random seed -> identical output on every run.

Usage:
    python -m src.generate_sample_data                # 300 files into data/sample
    python -m src.generate_sample_data --n 30 --out C:\\temp\\sample
"""
from __future__ import annotations

import argparse
import csv
import io
import os
import random
import shutil
import string
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from typing import Callable, Dict, List, Tuple

from faker import Faker

from . import get_logger, load_config, resolve_path
from .recognizers import verhoeff_generate_check_digit

log = get_logger("generate")

COMPANY = "Acme Bharat Pvt Ltd"
DOMAIN = "acmebharat.in"
PERSONAL_DOMAINS = ["gmail.com", "yahoo.co.in", "rediffmail.com", "outlook.com", "hotmail.com"]
BANKS = [("HDFC", "HDFC Bank"), ("ICIC", "ICICI Bank"), ("SBIN", "State Bank of India"),
         ("UTIB", "Axis Bank"), ("KKBK", "Kotak Mahindra Bank"), ("PUNB", "Punjab National Bank"),
         ("BARB", "Bank of Baroda"), ("CNRB", "Canara Bank"), ("YESB", "Yes Bank"), ("IDIB", "Indian Bank")]
UPI_HANDLES = ["okaxis", "okhdfcbank", "okicici", "oksbi", "ybl", "ibl", "paytm", "axl"]
CITIES = ["Mumbai", "Bengaluru", "Pune", "Hyderabad", "Chennai", "Gurugram", "Noida", "Kolkata", "Ahmedabad"]
HOSPITALS = ["City Care Hospital", "Sunrise Multispeciality Hospital", "Lotus Medical Centre", "Green Valley Hospital"]
CONDITIONS = [  # (phrase used in the note, follow-up phrase) - all contain health vocabulary
    ("was diagnosed with dengue and admitted to hospital", "The doctor has advised complete bed rest"),
    ("underwent appendicitis surgery last week", "A medical certificate from the surgeon is attached"),
    ("has been diagnosed with type 2 diabetes", "The prescription requires insulin twice a day"),
    ("is recovering from a fracture after a road accident", "Physiotherapy sessions are prescribed for 4 weeks"),
    ("is on maternity leave due to pregnancy complications", "The gynaecologist's prescription is attached"),
    ("was hospitalised for typhoid", "Blood test reports are attached with the medical certificate"),
    ("is undergoing treatment for hypertension", "Blood pressure readings must be monitored daily"),
    ("is being treated for a thyroid disorder", "The prescribed medication will continue for 3 months"),
    ("has been advised medical leave for clinical depression", "The psychiatrist has recommended therapy"),
    ("tested positive for malaria", "The diagnosis was confirmed by a blood test"),
]


# ---------------------------------------------------------------------------
# Identifier generators (all formats follow official rules; values are random)
# ---------------------------------------------------------------------------
class Gen:
    """Random but reproducible fake identifiers."""

    def __init__(self, seed: int):
        self.rng = random.Random(seed)
        self.fake = Faker("en_IN")
        self.fake.seed_instance(seed)

    # --- people -------------------------------------------------------------
    def person(self) -> Dict[str, str]:
        first, last = self.fake.first_name(), self.fake.last_name()
        return {"first": first, "last": last, "name": f"{first} {last}"}

    def email(self, p: Dict[str, str], work: bool = False) -> str:
        user = f"{p['first']}.{p['last']}".lower().replace(" ", "")
        if not work and self.rng.random() < 0.5:
            user += str(self.rng.randint(1, 99))
        return f"{user}@{DOMAIN if work else self.rng.choice(PERSONAL_DOMAINS)}"

    # --- Indian identifiers --------------------------------------------------
    def aadhaar(self, spaced: bool | None = None) -> str:
        """12 digits, first digit 2-9, last digit = Verhoeff check digit."""
        body = str(self.rng.randint(2, 9)) + "".join(str(self.rng.randint(0, 9)) for _ in range(10))
        num = body + verhoeff_generate_check_digit(body)
        spaced = self.rng.random() < 0.7 if spaced is None else spaced
        return f"{num[:4]} {num[4:8]} {num[8:]}" if spaced else num

    def pan(self, holder: str = "P", surname: str = "") -> str:
        """AAAAA9999A: 3 letters, holder type, surname/name initial, 4 digits, check letter."""
        letters = string.ascii_uppercase
        initial = surname[:1].upper() if surname[:1].isalpha() else self.rng.choice(letters)
        return ("".join(self.rng.choice(letters) for _ in range(3)) + holder + initial
                + f"{self.rng.randint(0, 9999):04d}" + self.rng.choice(letters))

    def bank(self) -> Tuple[str, str]:
        return self.rng.choice(BANKS)

    def ifsc(self, bank_code: str) -> str:
        """AAAA0XXXXXX: bank code, literal zero, 6-character branch code."""
        return f"{bank_code}0{self.rng.randint(0, 999999):06d}"

    def account(self) -> str:
        """11-16 digit account number (never starts with 0)."""
        length = self.rng.randint(11, 16)
        return str(self.rng.randint(1, 9)) + "".join(str(self.rng.randint(0, 9)) for _ in range(length - 1))

    def phone(self) -> str:
        num = str(self.rng.randint(6, 9)) + "".join(str(self.rng.randint(0, 9)) for _ in range(9))
        style = self.rng.randint(0, 3)
        return [f"+91 {num[:5]} {num[5:]}", f"+91-{num}", num, f"0{num}"][style]

    def upi(self, p: Dict[str, str]) -> str:
        return f"{p['first'].lower()}{self.rng.randint(10, 999)}@{self.rng.choice(UPI_HANDLES)}"

    def card(self) -> str:
        num = self.fake.credit_card_number(card_type="visa16")
        return " ".join(num[i:i + 4] for i in range(0, 16, 4))

    def ip(self) -> str:
        return f"10.{self.rng.randint(0, 255)}.{self.rng.randint(0, 255)}.{self.rng.randint(1, 254)}"

    def date(self) -> datetime:
        start = datetime(2021, 1, 1, tzinfo=timezone.utc)
        return start + timedelta(days=self.rng.randint(0, 1640), hours=self.rng.randint(8, 19),
                                 minutes=self.rng.randint(0, 59))

    def emp_id(self) -> str:
        return f"ABP{self.rng.randint(10000, 99999)}"

    def monthly_salary(self) -> int:
        return self.rng.randrange(28000, 260000, 500)


def inr(amount: int) -> str:
    """Indian digit grouping: 1234567 -> 12,34,567."""
    s = str(int(amount))
    last3, rest = s[-3:], s[:-3]
    groups = []
    while len(rest) > 2:
        groups.insert(0, rest[-2:])
        rest = rest[:-2]
    if rest:
        groups.insert(0, rest)
    return ",".join(groups + [last3]) if groups else last3


def money(g: Gen, amount: int) -> str:
    return g.rng.choice([f"₹{inr(amount)}", f"Rs. {inr(amount)}", f"INR {inr(amount)}"])


# ---------------------------------------------------------------------------
# Document = dict(kind, title, body lines/table, planted, tier, sender, subject, date)
# ---------------------------------------------------------------------------
def doc(tier: str, planted: set, subject: str, body: str = "", table: List[List] | None = None,
        sender: str = "", to: str = "", **extra) -> Dict:
    return {"tier": tier, "planted": planted, "subject": subject, "body": body, "table": table,
            "sender": sender, "to": to, **extra}


# ----- HIGH-RISK templates --------------------------------------------------
def t_payslip(g: Gen) -> Dict:
    p = g.person()
    code, bank_name = g.bank()
    basic = g.monthly_salary()
    hra, special = int(basic * 0.4), int(basic * 0.25)
    pf, tds = int(basic * 0.12), int(basic * 0.1)
    net = basic + hra + special - pf - tds
    month = g.rng.choice(["January", "February", "March", "April", "May", "June", "July", "August",
                          "September", "October", "November", "December"])
    lines = [
        f"{COMPANY} - Salary Slip for {month} {g.rng.randint(2021, 2025)}",
        "",
        f"Employee Name: {p['name']}",
        f"Employee ID: {g.emp_id()}",
        f"Designation: {g.rng.choice(['Software Engineer', 'Sales Manager', 'HR Executive', 'Analyst'])}",
        f"PAN: {g.pan('P', p['last'])}",
        f"Bank: {bank_name}",
        f"Bank Account No: {g.account()}",
        f"IFSC: {g.ifsc(code)}",
        "",
        "Earnings",
        f"Basic Salary: {money(g, basic)}",
        f"HRA: {money(g, hra)}",
        f"Special Allowance: {money(g, special)}",
        "Deductions",
        f"Provident Fund: {money(g, pf)}",
        f"TDS: {money(g, tds)}",
        f"Net Pay: {money(g, net)}",
        "",
        "This is a system generated payslip and does not require a signature.",
    ]
    return doc("High", {"PERSON", "IN_PAN", "BANK_ACCOUNT", "IN_IFSC", "SALARY"},
               f"Salary slip {month}", "\n".join(lines), sender=f"payroll@{DOMAIN}")


def t_payroll_sheet(g: Gen) -> Dict:
    rows = [["Employee Name", "Emp ID", "Aadhaar", "PAN", "Account No", "IFSC", "Net Salary (INR)", "Mobile"]]
    for _ in range(g.rng.randint(20, 200)):
        p = g.person()
        code, _ = g.bank()
        rows.append([p["name"], g.emp_id(), g.aadhaar(), g.pan("P", p["last"]), g.account(),
                     g.ifsc(code), g.monthly_salary(), g.phone()])
    return doc("High", {"PERSON", "IN_AADHAAR", "IN_PAN", "BANK_ACCOUNT", "IN_IFSC", "SALARY", "IN_PHONE"},
               "Payroll register", table=rows)


def t_offer_letter(g: Gen) -> Dict:
    p = g.person()
    ctc_lakh = g.rng.randint(6, 45)
    body = "\n".join([
        f"{COMPANY}",
        f"Date: {g.date():%d %B %Y}",
        "",
        f"Dear {p['name']},",
        "",
        f"We are pleased to offer you the position of {g.rng.choice(['Product Manager', 'Data Analyst', 'Senior Engineer', 'Account Executive'])} "
        f"at our {g.rng.choice(CITIES)} office.",
        f"Your annual CTC will be {money(g, ctc_lakh * 100000)} ({ctc_lakh} LPA), which includes a performance bonus "
        f"of up to {money(g, ctc_lakh * 8000)}.",
        f"Please confirm your acceptance by replying to {g.email(p)} or calling us back on the number you shared, "
        f"{g.phone()}.",
        "",
        "Your joining date and onboarding schedule will be shared by the HR team.",
        "",
        "Regards,",
        "Talent Acquisition Team",
    ])
    return doc("High", {"PERSON", "SALARY", "EMAIL_ADDRESS", "IN_PHONE"}, "Offer of employment", body,
               sender=f"careers@{DOMAIN}")


def t_medical_leave(g: Gen) -> Dict:
    p, manager = g.person(), g.person()
    cond, follow = g.rng.choice(CONDITIONS)
    days = g.rng.randint(3, 30)
    body = "\n".join([
        f"Hi {manager['first']},",
        "",
        f"Employee: {p['name']} ({g.emp_id()})",
        "",
        f"This is to inform you that the above employee {cond}. {follow}.",
        f"We have approved medical leave for {days} days. The discharge summary from "
        f"{g.rng.choice(HOSPITALS)} has been added to the HR file.",
        "",
        "Please treat this information as confidential.",
        "",
        "Thanks,",
        "HR Operations",
    ])
    return doc("High", {"PERSON", "HEALTH_INFO"}, "Medical leave approval", body,
               sender=f"hr@{DOMAIN}", to=g.email(manager, work=True))


def t_vendor_bank(g: Gen) -> Dict:
    p = g.person()
    code, bank_name = g.bank()
    vendor = f"{g.fake.last_name()} {g.rng.choice(['Traders', 'Logistics', 'Enterprises', 'Supplies'])}"
    body = "\n".join([
        "Dear Accounts Team,",
        "",
        f"Please update the bank details for our firm {vendor} for all future payments.",
        f"Beneficiary: {vendor}",
        f"Bank: {bank_name}",
        f"Account Number: {g.account()}",
        f"IFSC Code: {g.ifsc(code)}",
        f"UPI ID for small payments: {g.upi(p)}",
        f"Firm PAN: {g.pan('F')}",
        "",
        f"For any clarification contact {p['name']} on mobile {g.phone()}.",
        "",
        "Regards,",
        p["name"],
    ])
    return doc("High", {"PERSON", "BANK_ACCOUNT", "IN_IFSC", "IN_UPI", "IN_PAN", "IN_PHONE"},
               "Updated bank details for payments", body, sender=g.email(p))


def t_card_dispute(g: Gen) -> Dict:
    p = g.person()
    body = "\n".join([
        "Hello Support,",
        "",
        f"I was charged twice for my order. The charge appeared on my credit card {g.card()}.",
        f"Please refund the duplicate amount. You can reach me at {g.phone()} or {g.email(p)}.",
        "",
        "Thank you,",
        p["name"],
    ])
    return doc("High", {"PERSON", "CREDIT_CARD", "IN_PHONE", "EMAIL_ADDRESS"}, "Duplicate charge on my card",
               body, sender=g.email(p), to=f"support@{DOMAIN}")


def t_kyc_export(g: Gen) -> Dict:
    rows = [["Customer Name", "Aadhaar", "PAN", "Mobile", "Email", "City"]]
    for _ in range(g.rng.randint(30, 150)):
        p = g.person()
        rows.append([p["name"], g.aadhaar(), g.pan("P", p["last"]), g.phone(), g.email(p), g.rng.choice(CITIES)])
    return doc("High", {"PERSON", "IN_AADHAAR", "IN_PAN", "IN_PHONE", "EMAIL_ADDRESS"}, "Customer KYC export",
               table=rows)


# ----- MEDIUM-RISK templates ------------------------------------------------
def t_kyc_ticket(g: Gen) -> Dict:
    p = g.person()
    body = "\n".join([
        f"Ticket #{g.rng.randint(100000, 999999)} - KYC update request",
        "",
        f"Customer {p['name']} has requested an update to their KYC records.",
        f"Aadhaar number provided: {g.aadhaar()}",
        f"PAN provided: {g.pan('P', p['last'])}",
        f"Registered mobile: {g.phone()}",
        f"Email on file: {g.email(p)}",
        "",
        "Agent note: documents verified, please close the ticket after the update.",
    ])
    return doc("Medium", {"PERSON", "IN_AADHAAR", "IN_PAN", "IN_PHONE", "EMAIL_ADDRESS"}, "KYC update request",
               body, sender=f"support@{DOMAIN}")


def t_address_change(g: Gen) -> Dict:
    p = g.person()
    body = "\n".join([
        "Dear Customer Care,",
        "",
        f"I have moved to {g.rng.choice(CITIES)} and would like to update my address.",
        f"My Aadhaar number is {g.aadhaar()} for verification.",
        f"Old mobile: {g.phone()}. New mobile: {g.phone()}.",
        f"Please send the confirmation to {g.email(p)}.",
        "",
        "Regards,",
        p["name"],
    ])
    return doc("Medium", {"PERSON", "IN_AADHAAR", "IN_PHONE", "EMAIL_ADDRESS"}, "Address change request", body,
               sender=g.email(p), to=f"care@{DOMAIN}")


def t_contact_list(g: Gen) -> Dict:
    rows = [["Name", "Mobile", "Email", "Role"]]
    for _ in range(g.rng.randint(3, 5)):
        p = g.person()
        rows.append([p["name"], g.phone(), g.email(p), g.rng.choice(["Vendor", "Consultant", "Customer", "Partner"])])
    return doc("Medium", {"PERSON", "IN_PHONE", "EMAIL_ADDRESS"}, "Contact list", table=rows)


def t_background_check(g: Gen) -> Dict:
    p = g.person()
    body = "\n".join([
        "Background verification request",
        "",
        f"Candidate name: {p['name']}",
        f"PAN: {g.pan('P', p['last'])}",
        f"Aadhaar: {g.aadhaar()}",
        f"Contact number: {g.phone()}",
        "",
        "Please complete employment and education verification within 7 working days.",
    ])
    return doc("Medium", {"PERSON", "IN_PAN", "IN_AADHAAR", "IN_PHONE"}, "Background verification", body,
               sender=f"talent@{DOMAIN}")


def t_travel_booking(g: Gen) -> Dict:
    a, b = g.person(), g.person()
    body = "\n".join([
        "Hi Travel Desk,",
        "",
        f"Please book a hotel in {g.rng.choice(CITIES)} for two colleagues attending the client workshop.",
        f"Traveller 1: {a['name']}, mobile {g.phone()}, email {g.email(a, work=True)}",
        f"Traveller 2: {b['name']}, mobile {g.phone()}, email {g.email(b, work=True)}",
        f"The hotel needs a PAN for the invoice: {g.pan('P', a['last'])}",
        "",
        "Thanks",
    ])
    return doc("Medium", {"PERSON", "IN_PHONE", "EMAIL_ADDRESS", "IN_PAN"}, "Hotel booking for client workshop",
               body, sender=g.email(a, work=True), to=f"travel@{DOMAIN}")


def t_callback_list(g: Gen) -> Dict:
    lines = ["Customer callbacks pending for today:", ""]
    for i in range(3):
        p = g.person()
        lines.append(f"{i + 1}. {p['name']} - {g.phone()} - {g.email(p)} - wants a demo of the premium plan")
    lines += ["", "Please update the CRM after each call."]
    return doc("Medium", {"PERSON", "IN_PHONE", "EMAIL_ADDRESS"}, "Callback list", "\n".join(lines))


# ----- LOW-RISK templates ---------------------------------------------------
def t_meeting_invite(g: Gen) -> Dict:
    p = g.person()
    body = "\n".join([
        "Hi team,",
        "",
        f"{p['name']} from our implementation partner will join the quarterly review on Thursday.",
        f"Please share the agenda with them at {g.email(p, work=True)} before the meeting.",
        "",
        "Thanks",
    ])
    return doc("Low", {"PERSON", "EMAIL_ADDRESS"}, "Quarterly review - external attendee", body,
               sender=f"pmo@{DOMAIN}")


def t_vendor_contact(g: Gen) -> Dict:
    p = g.person()
    body = (f"Facility vendor contact for the {g.rng.choice(CITIES)} office: {p['name']}, "
            f"phone {g.phone()}. Call only for AC maintenance issues.")
    return doc("Low", {"PERSON", "IN_PHONE"}, "Facility vendor contact", body)


def t_it_ticket(g: Gen) -> Dict:
    p = g.person()
    body = "\n".join([
        f"IT ticket #{g.rng.randint(10000, 99999)}",
        f"Laptop of {p['name']} is unable to connect to the VPN.",
        f"Last seen IP address: {g.ip()}",
        "Resolution: reinstalled the VPN client and reset the certificate.",
    ])
    return doc("Low", {"PERSON", "IP_ADDRESS"}, "VPN connectivity issue", body, sender=f"itdesk@{DOMAIN}")


# ----- CLEAN templates (no personal data) -----------------------------------
PROJECTS = ["Phoenix CRM migration", "Q3 pricing refresh", "Warehouse automation", "Mobile app redesign",
            "Data lake modernisation", "Vendor portal", "Customer loyalty programme"]


def t_meeting_notes(g: Gen) -> Dict:
    proj = g.rng.choice(PROJECTS)
    body = "\n".join([
        f"Meeting notes - {proj}",
        "Attendees: Product, Engineering, Finance",
        "",
        "1. Sprint velocity is stable; two stories were carried over.",
        f"2. Budget utilisation is at {g.rng.randint(40, 95)} percent of the approved amount.",
        "3. Risks: dependency on the vendor API release and pending security review.",
        "4. Next steps: finalise the test plan and schedule a demo for the steering committee.",
    ])
    return doc("Clean", set(), f"Notes: {proj}", body)


def t_lunch_plan(g: Gen) -> Dict:
    place = g.rng.choice(["the rooftop cafe", "the South Indian place near the metro", "the new pizza place",
                          "the office cafeteria"])
    body = "\n".join([
        "Hi all,",
        "",
        f"Team lunch this Friday at {place}. The budget is Rs 600 per head, covered by the team fund.",
        "Reply with veg or non-veg preference so we can book a table.",
        "",
        "Cheers",
    ])
    return doc("Clean", set(), "Team lunch on Friday", body, sender=f"engagement@{DOMAIN}")


def t_project_update(g: Gen) -> Dict:
    proj = g.rng.choice(PROJECTS)
    body = "\n".join([
        f"Weekly status - {proj}",
        "",
        f"Overall status: {g.rng.choice(['Green', 'Amber', 'Green'])}",
        f"Completed: {g.rng.randint(5, 20)} of {g.rng.randint(21, 40)} planned tasks.",
        f"Purchase order PO-{g.rng.randint(4500000000, 4599999999)} for cloud credits has been approved.",
        f"Invoice value for the milestone is INR {inr(g.rng.randint(100000, 900000))}, due next month.",
        "Blockers: none. Next milestone is user acceptance testing.",
    ])
    return doc("Clean", set(), f"Status update: {proj}", body, sender=f"pmo@{DOMAIN}")


def t_policy(g: Gen) -> Dict:
    topic = g.rng.choice(["Work from home policy", "Travel reimbursement policy", "Information security policy",
                          "Leave and holiday policy"])
    body = "\n".join([
        f"{COMPANY} - {topic}",
        "",
        "1. Purpose: this policy sets out expectations for all employees.",
        "2. Scope: applies to all permanent and contract staff.",
        "3. Approvals must be recorded in the HR portal before the start of the activity.",
        "4. Exceptions require written approval from the department head.",
        "5. The policy is reviewed every year by the People team.",
    ])
    return doc("Clean", set(), topic, body)


def t_sales_sheet(g: Gen) -> Dict:
    rows = [["Region", "Quarter", "Units Sold", "Revenue (INR lakh)", "Target Achieved %"]]
    for region in ["North", "South", "East", "West"]:
        for q in ["Q1", "Q2", "Q3", "Q4"]:
            rows.append([region, q, g.rng.randint(200, 5000), g.rng.randint(20, 900), g.rng.randint(60, 130)])
    return doc("Clean", set(), "Regional sales summary", table=rows)


def t_release_notes(g: Gen) -> Dict:
    v = f"{g.rng.randint(1, 5)}.{g.rng.randint(0, 20)}.{g.rng.randint(0, 9)}"
    body = "\n".join([
        f"Release notes - version {v}",
        "- Improved dashboard loading time",
        "- Fixed an issue with CSV export encoding",
        "- Added dark mode to the settings page",
        "- Upgraded third-party libraries to patch known vulnerabilities",
    ])
    return doc("Clean", set(), f"Release {v}", body)


# Template, category, {extension: count at n=300}
TEMPLATES: List[Tuple[Callable[[Gen], Dict], str, Dict[str, int]]] = [
    # High (75)
    (t_payslip, "hr_salary_slip", {"docx": 9, "pdf": 5, "txt": 4}),
    (t_payroll_sheet, "payroll_spreadsheet", {"xlsx": 6, "csv": 6}),
    (t_offer_letter, "offer_letter", {"docx": 8, "pdf": 2}),
    (t_medical_leave, "medical_leave_note", {"eml": 8, "txt": 4, "docx": 3}),
    (t_vendor_bank, "vendor_bank_details", {"eml": 6, "pdf": 2}),
    (t_card_dispute, "card_dispute_ticket", {"eml": 5}),
    (t_kyc_export, "customer_kyc_export", {"csv": 4, "xlsx": 3}),
    # Medium (90)
    (t_kyc_ticket, "support_ticket_kyc", {"eml": 18, "txt": 12}),
    (t_address_change, "support_ticket_address", {"eml": 10, "txt": 5}),
    (t_contact_list, "contact_list", {"csv": 6, "xlsx": 4}),
    (t_background_check, "background_verification", {"docx": 5, "txt": 5, "eml": 5}),
    (t_travel_booking, "travel_booking", {"eml": 10}),
    (t_callback_list, "callback_list", {"txt": 10}),
    # Low (40)
    (t_meeting_invite, "meeting_invite", {"eml": 15, "txt": 5}),
    (t_vendor_contact, "vendor_contact", {"txt": 10}),
    (t_it_ticket, "it_ticket", {"eml": 10}),
    # Clean (95)
    (t_meeting_notes, "meeting_notes", {"txt": 15, "docx": 10}),
    (t_lunch_plan, "lunch_plan", {"eml": 15}),
    (t_project_update, "project_update", {"eml": 15, "pdf": 3}),
    (t_policy, "policy_document", {"docx": 5, "pdf": 2}),
    (t_sales_sheet, "sales_summary", {"csv": 5, "xlsx": 5}),
    (t_release_notes, "release_notes", {"txt": 20}),
]


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------
def _text_of(d: Dict) -> str:
    if d["table"] is not None:
        return "\n".join(", ".join(str(c) for c in row) for row in d["table"])
    return d["body"]


def write_eml(path: Path, d: Dict, g: Gen, when: datetime) -> None:
    msg = EmailMessage()
    msg["From"] = d["sender"] or f"noreply@{DOMAIN}"
    msg["To"] = d["to"] or f"team@{DOMAIN}"
    msg["Subject"] = d["subject"]
    msg["Date"] = format_datetime(when)
    msg["Message-ID"] = f"<{g.rng.randint(10**9, 10**10)}@{DOMAIN}>"
    msg.set_content(_text_of(d))
    path.write_bytes(bytes(msg))


def write_txt(path: Path, d: Dict, *_):
    path.write_text(_text_of(d), encoding="utf-8")


def write_csv(path: Path, d: Dict, *_):
    table = d["table"] or [["Note"], [d["body"]]]
    with path.open("w", encoding="utf-8", newline="") as fh:
        csv.writer(fh).writerows(table)


def write_xlsx(path: Path, d: Dict, *_):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = d["subject"][:30]
    for row in d["table"] or [["Note"], [d["body"]]]:
        ws.append(row)
    wb.save(path)


def write_docx(path: Path, d: Dict, g: Gen, when: datetime):
    import docx

    document = docx.Document()
    lines = _text_of(d).split("\n")
    document.add_heading(lines[0] if lines else d["subject"], level=1)
    for line in lines[1:]:
        document.add_paragraph(line)
    document.core_properties.created = when.replace(tzinfo=None)
    document.core_properties.author = COMPANY
    document.save(path)


def write_pdf(path: Path, d: Dict, g: Gen, when: datetime):
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    width, height = A4
    y = height - 60
    for line in _text_of(d).split("\n"):
        # Built-in Helvetica has no rupee glyph, so PDFs spell the currency out.
        line = line.replace("₹", "Rs. ")
        while pdfmetrics.stringWidth(line, "Helvetica", 10) > width - 100:
            cut = line.rfind(" ", 0, 95)
            cut = cut if cut > 0 else 95
            c.setFont("Helvetica", 10)
            c.drawString(50, y, line[:cut])
            line, y = line[cut:].lstrip(), y - 14
        c.setFont("Helvetica", 10)
        c.drawString(50, y, line)
        y -= 14
        if y < 60:
            c.showPage()
            y = height - 60
    c.save()
    buf.seek(0)
    writer = PdfWriter(clone_from=PdfReader(buf))
    writer.add_metadata({"/CreationDate": when.strftime("D:%Y%m%d%H%M%S+00'00'"), "/Producer": COMPANY,
                         "/Title": d["subject"]})
    with path.open("wb") as fh:
        writer.write(fh)


WRITERS = {"eml": write_eml, "txt": write_txt, "csv": write_csv, "xlsx": write_xlsx,
           "docx": write_docx, "pdf": write_pdf}


def _prepare_output(out: Path) -> None:
    """Empty the output folder, but only if it is empty or was generated by us."""
    if out.exists() and any(out.iterdir()):
        if not (out / "ground_truth.csv").exists():
            raise SystemExit(f"Refusing to overwrite {out}: it has files but no ground_truth.csv "
                             "(does not look like generated sample data).")
        for child in out.iterdir():
            if child.name == ".gitkeep":
                continue
            shutil.rmtree(child) if child.is_dir() else child.unlink()
    out.mkdir(parents=True, exist_ok=True)


def generate(out_dir: Path, n: int = 300, seed: int = 42) -> Dict[str, int]:
    """Generate ~n synthetic files plus ground_truth.csv into ``out_dir``."""
    _prepare_output(out_dir)
    g = Gen(seed)
    scale = n / 300.0
    plan = []
    for template, category, ext_counts in TEMPLATES:
        for ext, count in ext_counts.items():
            plan.extend([(template, category, ext)] * max(1, round(count * scale)))
    g.rng.shuffle(plan)

    truth_rows, tiers = [], {}
    for i, (template, category, ext) in enumerate(plan, start=1):
        d = template(g)
        when = g.date()
        name = f"{category}_{i:04d}.{ext}"
        # Put files into department sub-folders so the recursive walk is exercised.
        dept = ("hr" if category in {"hr_salary_slip", "payroll_spreadsheet", "offer_letter", "medical_leave_note",
                                     "background_verification", "policy_document"}
                else "support" if category.startswith(("support", "card", "callback", "customer"))
                else "finance" if category in {"vendor_bank_details", "sales_summary"}
                else "general")
        path = out_dir / dept / name
        path.parent.mkdir(parents=True, exist_ok=True)
        WRITERS[ext](path, d, g, when)
        ts = when.timestamp()
        os.utime(path, (ts, ts))
        truth_rows.append({"file": str(path.relative_to(out_dir)).replace("\\", "/"), "category": category,
                           "file_type": ext, "expected_risk": d["tier"],
                           "planted_entities": ";".join(sorted(d["planted"]))})
        tiers[d["tier"]] = tiers.get(d["tier"], 0) + 1

    # One deliberately corrupt file: demonstrates graceful skipping.
    (out_dir / "general" / "corrupt_scan_0000.pdf").write_bytes(b"%PDF-1.4\n\x00\x9c this is not a valid pdf \xff")

    with (out_dir / "ground_truth.csv").open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["file", "category", "file_type", "expected_risk", "planted_entities"])
        w.writeheader()
        w.writerows(truth_rows)
    (out_dir / "README_SYNTHETIC_DATA.txt").write_text(
        "All files in this folder are SYNTHETIC and generated by src/generate_sample_data.py.\n"
        "No real person's data is included. Aadhaar/PAN/IFSC values follow official formats only.\n",
        encoding="utf-8")
    log.info("Generated %d files in %s: %s", len(plan), out_dir, tiers)
    return {"files": len(plan), **tiers}


def main(argv=None) -> None:
    cfg = load_config()
    ap = argparse.ArgumentParser(description="Generate synthetic Indian corporate data (all fake).")
    ap.add_argument("--n", type=int, default=300, help="approximate number of files (default 300)")
    ap.add_argument("--out", default=cfg["paths"]["sample_dir"], help="output folder (default data/sample)")
    ap.add_argument("--seed", type=int, default=42, help="random seed (default 42)")
    args = ap.parse_args(argv)
    out = resolve_path(args.out)
    stats = generate(out, args.n, args.seed)
    print(f"Created {stats['files']} synthetic files in {out}")
    print("  " + ", ".join(f"{k}: {v}" for k, v in stats.items() if k != "files"))
    print(f"  Ground truth: {out / 'ground_truth.csv'}")


if __name__ == "__main__":
    main()
